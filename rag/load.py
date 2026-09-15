"""Extract text from supported document formats."""

from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {
    ".txt", ".md", ".rst", ".csv", ".tsv", ".json", ".jsonl",
    ".yaml", ".yml", ".xml", ".html", ".htm", ".pdf", ".docx",
}


def load_document(path: Path) -> str:
    """Read a document as text.

    Args:
        path: Path to a supported document.

    Returns:
        Extracted text, with surrounding whitespace removed.

    Raises:
        ValueError: The file format is unsupported.
    """
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file format: {suffix or '(no extension)'}")

    if suffix == ".pdf":
        text = "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    elif suffix == ".docx":
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(parts)
    elif suffix in {".html", ".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8-sig"), "html.parser")
        for element in soup(["script", "style"]):
            element.decompose()
        text = soup.get_text(separator="\n", strip=True)
    else:
        text = path.read_text(encoding="utf-8-sig")

    return text.replace("\x00", "").strip()
